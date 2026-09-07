import pathlib,sys,json,importlib.metadata
from docx import Document
from pptx import Presentation
from openpyxl import Workbook,load_workbook
from reportlab.pdfgen import canvas
from pypdf import PdfReader
import pdfplumber
import pypdfium2
from PIL import Image
out=pathlib.Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
doc=Document();doc.add_paragraph('Nodex workspace runtime');doc.save(out/'sample.docx')
assert Document(out/'sample.docx').paragraphs[0].text=='Nodex workspace runtime'
ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='Nodex';ppt.save(out/'sample.pptx')
assert Presentation(out/'sample.pptx').slides[0].shapes.title.text=='Nodex'
wb=Workbook();wb.active['A1']=21;wb.active['B1']='=A1*2';wb.save(out/'sample.xlsx')
assert load_workbook(out/'sample.xlsx').active['B1'].value=='=A1*2'
pdf=canvas.Canvas(str(out/'sample.pdf'));pdf.drawString(72,720,'Nodex workspace runtime');pdf.save()
assert 'Nodex workspace runtime' in PdfReader(out/'sample.pdf').pages[0].extract_text()
with pdfplumber.open(out/'sample.pdf') as document: assert 'Nodex' in document.pages[0].extract_text()
with pypdfium2.PdfDocument(out/'sample.pdf') as document: document[0].render().to_pil().save(out/'sample.png')
assert Image.open(out/'sample.png').width>0
print(json.dumps({'python':sys.version.split()[0],'files':sorted(p.name for p in out.iterdir())}))
